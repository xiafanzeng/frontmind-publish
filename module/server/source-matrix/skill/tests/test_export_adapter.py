"""Real workbook regressions for the export-research boundary."""
import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from scripts.export_adapter import read_export, research_matrix

HEAD = ['引用日期', '模型', '监控问题', '内容标题', '内容链接', '媒体名称', '媒体url']
ROW = ['2026-10-22', '元宝-手机版', '如何选软件测试服务', '测试资质说明', 'https://a.example/p?id=1', '甲媒体', 'a.example']
BRIEF = {'brand_product': '', 'target_industry': '软件测试服务', 'target_questions': ['如何选软件测试服务'], 'regions': None}

class ExportAdapterTests(unittest.TestCase):
    def workbook(self, rows=None, summaries=True, detail=True, headers=None):
        folder = tempfile.TemporaryDirectory(); self.addCleanup(folder.cleanup)
        path = Path(folder.name) / 'export.xlsx'
        w = Workbook(); w.remove(w.active)
        if detail:
            s = w.create_sheet('任意明细名'); s.append(headers or HEAD)
            for row in (rows if rows is not None else [ROW, ROW]): s.append(row)
        if summaries:
            s = w.create_sheet('媒体统计'); s.append(['媒体名称','媒体URL','引用日期','引用次数','引用占比']); s.append(['甲媒体','a.example','2026-10-16 — 2026-10-22',2,'100%'])
            s = w.create_sheet('内容统计'); s.append(['内容标题','内容URL','媒体名称','媒体url','引用次数']); s.append(['测试资质说明', ROW[4], '甲媒体','a.example',2])
        w.save(path); w.close(); return path

    def test_duplicates_retained_and_three_views_not_added(self):
        a = read_export(self.workbook())
        self.assertEqual(a['metrics']['raw_rows'], 2)
        self.assertEqual(a['metrics']['distinct_export_rows'], 1)
        self.assertEqual(a['metrics']['extra_identical_rows'], 1)
        self.assertEqual(a['metrics']['unique_urls'], 1)
        self.assertEqual(a['reconciliation']['media_reported_total'], 2)
        self.assertEqual(a['reconciliation']['content_reported_total'], 2)
        self.assertEqual(len(a['records']), 2)
        self.assertIsNone(a['metrics']['exact_citation_count'])
        self.assertNotIn('answer_record_id', a['records'][0])
        self.assertNotIn('evidence_id', a['records'][0])
        self.assertNotIn('published_at', a['records'][0])

    def test_source_locator_points_to_real_row(self):
        a = read_export(self.workbook())
        self.assertEqual(a['records'][0]['locator']['row'], 2)
        self.assertEqual(a['records'][1]['locator']['row'], 3)
        self.assertNotEqual(a['records'][0]['locator_key'], a['records'][1]['locator_key'])

    def test_header_aliases_reordered_and_unknown_columns_preserved(self):
        a = read_export(self.workbook(rows=[['extra', *reversed(ROW)]], headers=['备注', *reversed(HEAD)]))
        self.assertEqual(a['records'][0]['content_url'], ROW[4])
        self.assertEqual(a['records'][0]['raw']['备注'], 'extra')

    def test_summary_only_is_not_zero_detail_or_industry(self):
        a = read_export(self.workbook(detail=False))
        self.assertEqual(a['status'], 'aggregate_only')
        self.assertIsNone(a['metrics']['raw_rows'])
        self.assertIsNone(a['metrics']['unique_urls'])
        self.assertIsNone(a['metrics']['exact_citation_count'])
        r = research_matrix(a, BRIEF)
        self.assertEqual(r['mode'], 'export_research')
        self.assertFalse(r['industry_distribution_available'])
        self.assertEqual(r['budget_plans'], [])

    def test_summary_only_offers_traceable_groups_for_agent_judgment(self):
        a=read_export(self.workbook(detail=False))
        self.assertEqual(len(a['groups']),1)
        self.assertEqual(a['groups'][0]['reported_count'],2)
        self.assertIsNone(a['groups'][0]['raw_rows'])
        self.assertTrue(a['groups'][0]['locator_keys'])

    def test_content_only_preserves_limited_summary_candidate_without_fake_detail(self):
        path=self.workbook(detail=False); w=load_workbook(path); w.remove(w['媒体统计']);w.save(path);w.close()
        a=read_export(path)
        self.assertEqual(len(a['groups']),1)
        self.assertEqual(a['groups'][0]['reported_count'],2)
        self.assertIsNone(a['metrics']['raw_rows'])
        d=self.decision_from_summary(a)
        r=research_matrix(a,BRIEF,d)
        self.assertEqual(r['status'],'research_draft')
        self.assertIsNone(r['matrix'][0]['raw_rows'])

    def decision_from_summary(self,a):
        g=a['groups'][0]
        return {'input_version':research_matrix(a,BRIEF)['input_version'],
                'items':[{'candidate_key':g['candidate_key'],'primary_role':'industry','matched_questions':BRIEF['target_questions'],
                          'reason':'仅按原汇总内容主题研究，不声称逐条引用命中。','content_direction':'服务选型说明',
                          'references':g['locator_keys']}], 'drafts':[]}

    def test_empty_assumed_region_does_not_authorize_regional_assignment(self):
        a=read_export(self.workbook());d=self.decision(a)
        d['items'][0].update(primary_role='regional',verification_items=['确认覆盖地区'])
        self.assertEqual(research_matrix(a,dict(BRIEF,regions=[]),d)['status'],'invalid_decisions')

    def test_drift_reported_instead_of_adjusting_detail(self):
        a = read_export(self.workbook(rows=[ROW]))
        self.assertEqual(a['metrics']['raw_rows'], 1)
        self.assertEqual(a['reconciliation']['status'], 'mismatch')

    def test_same_domain_different_names_not_merged(self):
        other = list(ROW); other[5] = '乙媒体'
        a = read_export(self.workbook(rows=[ROW, other]))
        self.assertEqual(len(a['groups']), 2)

    def test_model_filter_recalculates_rows_without_touching_original_summaries(self):
        other = list(ROW); other[1] = '豆包-手机版'; other[0]='2026-10-21'; other[2]='另一监控问题'
        a = read_export(self.workbook(rows=[ROW, other]))
        r = research_matrix(a, dict(BRIEF, sample_models=['元宝-手机版']))
        self.assertEqual(r['observation']['metrics']['raw_rows'], 1)
        self.assertEqual(r['observation']['reconciliation']['status'], 'not_comparable_filtered')
        self.assertEqual(r['observation']['reconciliation']['detail_rows'],1)
        self.assertEqual(r['observation']['observed_dates'],{'2026-10-22':1})
        self.assertEqual(r['observation']['monitored_questions'],{'如何选软件测试服务':1})
        self.assertEqual(a['metrics']['raw_rows'], 2)

    def test_missing_industry_keeps_observation_but_does_not_invent_brief(self):
        r = research_matrix(read_export(self.workbook()), {})
        self.assertEqual(r['status'], 'needs_requirements')
        self.assertIsNone(r['brief']['target_industry'])
        self.assertIsNone(r['brief']['regions'])

    def test_new_goal_invalidates_saved_semantic_decisions_but_budget_does_not(self):
        a = read_export(self.workbook()); initial = research_matrix(a, BRIEF)
        d = {'input_version': initial['input_version'], 'items': [], 'drafts': []}
        r = research_matrix(a, dict(BRIEF, budget={'total_cap':'100'}), d)
        self.assertNotEqual(r['status'], 'stale_decisions')
        r = research_matrix(a, dict(BRIEF, audience='采购经理'), d)
        self.assertEqual(r['status'], 'stale_decisions')
        self.assertEqual(r['matrix'], [])

    def decision(self, a):
        r = research_matrix(a, BRIEF); g = a['groups'][0]
        return {'input_version':r['input_version'], 'items':[{'candidate_key':g['candidate_key'], 'primary_role':'industry', 'matched_questions':BRIEF['target_questions'], 'reason':'内容主题与测试服务选型问题相符，账号与报价待确认。', 'content_direction':'按项目类型解释报告用途与交付范围', 'references':[a['records'][0]['locator_key']]}], 'drafts':[]}

    def test_decision_is_research_only_and_has_no_purchasable_identity(self):
        a = read_export(self.workbook()); r = research_matrix(a, BRIEF, self.decision(a))
        self.assertEqual(r['status'], 'research_draft')
        self.assertEqual(r['matrix'][0]['channel_status'], 'research_only')
        self.assertIsNone(r['matrix'][0]['resource_id'])
        self.assertIsNone(r['ready_to_publish'])
        self.assertFalse(r['execution_recommended'])

    def test_media_introduction_is_separate_from_content_direction(self):
        a = read_export(self.workbook()); d = self.decision(a)
        intro = '以技术文章介绍软件服务的内容平台。'
        d['items'][0]['media_intro'] = intro
        r = research_matrix(a, BRIEF, d)
        self.assertEqual(r['status'], 'research_draft')
        self.assertEqual(r['matrix'][0]['media_intro'], intro)
        self.assertEqual(r['matrix'][0]['content_direction'], d['items'][0]['content_direction'])
        self.assertEqual(r['input_version'], d['input_version'])
        for invalid in (None, [], '   '):
            with self.subTest(invalid=invalid):
                d['items'][0]['media_intro'] = invalid
                self.assertEqual(research_matrix(a, BRIEF, d)['status'], 'invalid_decisions')

    def test_foreign_reference_rejected(self):
        a = read_export(self.workbook()); d = self.decision(a); d['items'][0]['references'] = ['invented-row']
        self.assertEqual(research_matrix(a, BRIEF, d)['status'], 'invalid_decisions')

    def ranked_decision(self):
        other = list(ROW); other[4:7] = ['https://b.example/post', '乙媒体', 'b.example']
        a = read_export(self.workbook(rows=[ROW, other]))
        d = {'input_version':research_matrix(a, BRIEF)['input_version'], 'items':[], 'drafts':[]}
        for index, group in enumerate(a['groups']):
            d['items'].append({'candidate_key':group['candidate_key'], 'primary_role':'industry',
                'matched_questions':BRIEF['target_questions'], 'reason':'软件测试主题与选型问题相关。',
                'content_direction':'解释项目范围与交付材料。', 'references':[group['locator_keys'][0]],
                'layer_rank':2-index, 'ranking_reason':'按目标问题匹配及内容分工确定同层次序。'})
        return a, d

    def test_layer_ranking_is_explicit_and_budget_independent(self):
        a, d = self.ranked_decision()
        r = research_matrix(a, BRIEF, d)
        self.assertEqual(r['status'], 'research_draft')
        self.assertEqual(sorted(x['layer_rank'] for x in r['matrix']), [1, 2])
        budget = research_matrix(a, dict(BRIEF, budget={'total_cap':'100'}), d)
        self.assertEqual(r['matrix'], budget['matrix'])
        self.assertEqual(r['input_version'], budget['input_version'])

    def test_layer_ranking_rejects_invalid_or_incomplete_orders(self):
        a, original = self.ranked_decision()
        variants = []
        for value in (True, 0, -1, 1.5, '1'):
            d=copy.deepcopy(original);d['items'][0]['layer_rank']=value;variants.append(d)
        d=copy.deepcopy(original);d['items'][0]['layer_rank']=1;variants.append(d)
        d=copy.deepcopy(original);d['items'][0]['layer_rank']=3;variants.append(d)
        d=copy.deepcopy(original);del d['items'][1]['layer_rank'];variants.append(d)
        d=copy.deepcopy(original);d['items'][0]['ranking_reason']=' ';variants.append(d)
        for d in variants:
            with self.subTest(items=d['items']):
                self.assertEqual(research_matrix(a, BRIEF, d)['status'], 'invalid_decisions')

    def test_unknown_question_and_unjustified_regional_role_rejected(self):
        a = read_export(self.workbook()); d = self.decision(a); d['items'][0]['matched_questions'] = ['invented question']
        self.assertEqual(research_matrix(a, BRIEF, d)['status'], 'invalid_decisions')
        d = self.decision(a); d['items'][0]['primary_role'] = 'regional'
        self.assertEqual(research_matrix(a, BRIEF, d)['status'], 'invalid_decisions')

    def test_formula_preserved_as_untrusted_text_never_evaluated(self):
        row = list(ROW); row[3] = '=HYPERLINK("https://example.com","run")'
        a = read_export(self.workbook(rows=[row]))
        self.assertTrue(a['records'][0]['title'].startswith('=HYPERLINK'))

    def test_wrong_selected_model_gives_empty_filtered_observation(self):
        a = read_export(self.workbook()); r = research_matrix(a, dict(BRIEF, sample_models=['不存在的模型']))
        self.assertEqual(r['observation']['metrics']['raw_rows'], 0)
        self.assertEqual(r['status'], 'empty_selection')

    def test_malformed_decisions_return_structured_errors_instead_of_crashing_preview(self):
        a=read_export(self.workbook())
        for field,value in [('references',None),('verification_items','确认身份'),('matched_questions','非数组'),('candidate_key',[])]:
            with self.subTest(field=field):
                d=self.decision(a);d['items'][0][field]=value
                self.assertEqual(research_matrix(a,BRIEF,d)['status'],'invalid_decisions')
        d=self.decision(a);d['drafts']=[None]
        self.assertEqual(research_matrix(a,BRIEF,d)['status'],'invalid_decisions')
        self.assertEqual(research_matrix(a,BRIEF,[])['status'],'invalid_decisions')

if __name__ == '__main__': unittest.main()
